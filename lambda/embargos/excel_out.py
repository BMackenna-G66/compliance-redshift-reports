"""Escritura del Excel de salida (datos puros, sin formulas)."""
from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FUENTE = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name=FUENTE, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FUENTE, size=10)
ANCHOS = {"nombre_completo": 38, "numeros_proceso": 34, "numero_proceso": 26,
          "flags": 46, "archivo_origen": 30, "hojas_origen": 26, "demandante": 30,
          "nombre_en_sistema": 28, "numeros_oficio": 22}


def _hoja(wb: Workbook, titulo: str, filas: List[Dict], columnas: Optional[List[str]] = None):
    ws = wb.create_sheet(titulo[:31])
    if not filas:
        ws["A1"] = "Sin registros"
        ws["A1"].font = BODY_FONT
        return ws
    columnas = columnas or list(filas[0].keys())
    ws.append(columnas)
    for c in range(1, len(columnas) + 1):
        celda = ws.cell(row=1, column=c)
        celda.font, celda.fill = HEADER_FONT, HEADER_FILL
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    for fila in filas:
        ws.append([fila.get(c, "") for c in columnas])
    for c, nombre in enumerate(columnas, start=1):
        ws.column_dimensions[get_column_letter(c)].width = ANCHOS.get(nombre, 18)
    for row in ws.iter_rows(min_row=2):
        for celda in row:
            celda.font = BODY_FONT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{ws.max_row}"
    return ws


def escribir_excel(destino: Path,
                   clientes: List[Dict],
                   no_clientes: List[Dict],
                   descartados: List[Dict],
                   resumen: List[Dict]) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Resumen")
    ws.append(["Indicador", "Valor"])
    for c in ("A1", "B1"):
        ws[c].font, ws[c].fill = HEADER_FONT, HEADER_FILL
    for item in resumen:
        ws.append([item["indicador"], item["valor"]])
    for row in ws.iter_rows(min_row=2):
        for celda in row:
            celda.font = BODY_FONT
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 60

    _hoja(wb, "Clientes", clientes)
    _hoja(wb, "No clientes", no_clientes)
    _hoja(wb, "Descartados y revision", descartados)

    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino
